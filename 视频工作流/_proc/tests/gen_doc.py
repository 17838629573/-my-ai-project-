"""生成 35 项推进文档（xlsx）
契约: tests/gen_doc
  输入: tests/cases.CASES、tests/harness.CRIT、motion.beat.CAP
  输出: 推进路线.xlsx（4 表）
  约束: 批次与搜索关键词来自本文件 BATCH 表，不就地拍
"""
import os, sys, json
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT); sys.path.insert(0, HERE)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from cases import CASES
import harness as H
import motion.beat as B

# 批次 -> (目标, 能力, 搜索关键词)
BATCH = {
 0: ("修 A1 走路：脚穿地/浮空", [], "walk gait foot contact penetration"),
 1: ("单体动作 A4/A5", ["turn","wave"], "turn in place 180 head torso foot sequence; wave gesture shoulder elbow wrist"),
 2: ("跑跳蹲 A2/A3/A7", ["run","brake","jump","crouch"], "run cycle braking deceleration; jump squat takeoff landing; crouch squat down"),
 3: ("抓取 B8/B9（能力已齐，写执行）", [], "grasp release hand object attachment"),
 4: ("刚体物理 C 组", ["rigid_body","ball","bounce","ramp","stack","pendulum","toppling"], "rigid body impulse restitution friction; sphere rolling ramp; box stacking stability; pendulum momentum transfer; domino toppling"),
 5: ("投掷/踢/推/门 B 组", ["throw","catch","kick","push","hinge_door","rope"], "throw projectile motion release; catch prediction; kick ball; push box; hinge door rotation; rope tension constraint"),
 6: ("多主体 D 组", ["multi_actor","contact","crowd_collide","quadruped"], "multi actor contact high five; crowd collision avoidance; quadruped walk cycle"),
 7: ("复杂序列 E 组", ["reach_grab","roll","carry_box","climb"], "running jump grab roll; carry box; ladder climb hand alternation"),
 8: ("压力测试 F 组", ["ccd","broadphase","gravity_off","friction","overlap_resolve"], "continuous collision detection tunneling; broadphase 50 bodies; zero gravity; high friction; overlap resolution"),
}
# 用例 -> 批次
CASE_BATCH = {
 "A1":0,
 "A4":1,"A5":1,
 "A2":2,"A3":2,"A7":2,
 "B8":3,"B9":3,
 "C18":4,"C20":4,"C17":4,"C16":4,"C19":4,"C15":4,
 "B10":5,"B11":5,"A6":5,"B12":5,"B14":5,"B13":5,
 "D21":6,"D22":6,"D23":6,"D24":6,
 "E25":7,"E26":7,"E27":7,"E28":7,"E29":7,
 "F30":8,"F31":8,"F32":8,"F33":8,"F34":8,"F35":8,
}

wb = Workbook()
HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", start_color="3A6EA5")
WRAP = Alignment(wrap_text=True, vertical="top")

def sheet(title, head, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in range(1, len(head)+1):
        ws.cell(1, c).font = HDR; ws.cell(1, c).fill = FILL
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = WRAP
    ws.freeze_panes = "A2"
    return ws

# ---- 表1 推进路线
rows = []
for cid, grp, name, caps, crits, seed in CASES:
    miss = [c for c in caps if c not in B.CAP]
    st = "PASS/FAIL待测" if not miss else ("缺能力" if miss == caps else "部分缺")
    rows.append([CASE_BATCH.get(cid, 9), cid, grp, name, ",".join(caps),
                 ",".join(miss) or "-", ",".join(crits), seed, st])
rows.sort(key=lambda r: (r[0], r[1]))
sheet("推进路线", ["批次","ID","组","用例","需求能力","缺失能力","判据","种子","状态"], rows,
      [6,6,5,34,26,26,30,7,14])

# ---- 表2 能力清单
capmap = {}
for cid, grp, name, caps, crits, seed in CASES:
    for c in caps:
        capmap.setdefault(c, {"cases": [], "miss": c not in B.CAP})
        capmap[c]["cases"].append(cid)
crows = []
for c, v in sorted(capmap.items()):
    b = None
    for bi, (_, cl, _) in BATCH.items():
        if c in cl: b = bi
    hint = ""
    for bi, (_, cl, k) in BATCH.items():
        if c in cl: hint = k
    crows.append([c, "已有" if not v["miss"] else "缺失", b if b is not None else "-",
                  ",".join(v["cases"]), hint])
crows.sort(key=lambda r: (r[1] != "缺失", r[2] if isinstance(r[2], int) else 99))
sheet("能力清单", ["能力","状态","批次","被引用用例","搜索关键词(用到才搜)"], crows, [18,7,6,30,46])

# ---- 表3 判据阈值
krows = [[k, v[0], v[1], v[2]] for k, v in H.CRIT.items()]
sheet("判据阈值", ["判据","阈值","比较","出处"], krows, [18,10,7,80])

# ---- 表4 批次计划
brows = []
for bi in sorted(BATCH):
    tgt, caps, hint = BATCH[bi]
    ids = sorted([c for c, b in CASE_BATCH.items() if b == bi])
    brows.append([bi, tgt, ",".join(ids), len(ids), ",".join(caps) or "-", hint])
sheet("批次计划", ["批次","目标","用例","项数","新增能力","搜索关键词"], brows, [6,28,26,6,34,52])

del wb["Sheet"]
out = os.path.join(ROOT, "推进路线.xlsx")
wb.save(out)
print("已生成", out, "表:", wb.sheetnames)
