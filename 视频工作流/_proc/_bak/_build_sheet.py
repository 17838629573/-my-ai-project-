# -*- coding: utf-8 -*-
"""把 FORMULA_REGISTRY 的全部缺口项建成工作表，并按施工批次排序。"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

REG = 'FORMULA_REGISTRY.json'
d = json.load(open(REG, encoding='utf-8'))

WAVE = {
    1: 'W1 地基：动作时间线',
    2: 'W2 动作能力',
    3: 'W3 相机运动',
    4: 'W4 场景与陈设',
    5: 'W5 动态元素',
    6: 'W6 人物外观细节',
    7: 'W7 音频',
    8: 'W8 形体与后处理',
}

# 项 -> (批次, 依赖, 搜索关键词, 验证方式)
M = {
 'prompt.action_beat_timeline': (1,'—','ADAPT capability function animation timeline scheduler; Behavior Markup Language BML scheduling','多段动作串成 12s，段边界无跳变'),
 'prompt.action_sit':           (2,'beat_timeline, char.ik','sit down animation IK contact constraint biped; sit to stand STS phase','臀贴凳 脚踩地，支撑面内'),
 'prompt.action_stand_up':      (2,'beat_timeline, char.ik','stand up animation center of mass transfer STS four phase','重心前移→离凳→站直'),
 'prompt.action_reach_grab':    (2,'beat_timeline','reach grasp animation inverse kinematics CCD FABRIK limb','手到目标点，肘方向合理'),
 'prompt.action_carry_prop':    (2,'reach_grab','prop handoff ownership attachment animation ADAPT PropHandoff','道具跟随手，不穿模'),
 'prompt.action_prop_release':  (2,'carry_prop','prop release placement animation ownership transfer','道具脱离手，落到面'),
 'prompt.action_gaze_shift':    (2,'beat_timeline','gaze shift head eye movement saccade amplitude formula','视线转向目标，头眼时序'),
 'prompt.action_expression':    (2,'beat_timeline','facial expression animation blendshape smile mouth corner','嘴角上扬可测'),
 'prompt.action_legs_cross':    (2,'sit','crossed legs sitting pose rig IK','双腿交叠不穿模'),
 'prompt.action_page_flip':     (2,'carry_prop, dyn.paper_bend','page flip paper deformation animation curve','纸面卷曲再落下'),
 'prompt.action_finger_tap':    (2,'手LOD1','finger tap animation hand rig joint chain','指尖上下，掌不动'),
 'char.ik':                     (2,'—','inverse kinematics contact constraint biped animation solver','IK 残差 <1e-6，关节不超限'),

 'prompt.cam_shot_size':        (3,'cam.projection','shot size framing MLS MS MCU subject height ratio','景别切换人物占画面比例达标'),
 'prompt.cam_pan':              (3,'cam.projection','camera pan angular velocity formula horizontal','横摇角速度恒定无跳'),
 'prompt.cam_dolly_in':         (3,'cam.projection','camera dolly zoom perspective interpolation formula','推近时透视缩放单调递增'),
 'prompt.cam_tracking':         (3,'cam_pan, cam_dolly_in','camera tracking follow smoothing deadzone formula','跟拍主体保持在框内'),
 'cam.parallax':                (3,'cam.px_per_m','2D parallax layer scroll rate depth formula','各层速率比 = 深度反比'),

 'prompt.env_indoor_pack':      (4,'—','procedural interior room generation shape grammar; procedural cafe layout','室内包可出，层序正确'),
 'prompt.env_props':            (4,'env_indoor_pack','procedural furniture chair table generation parameter','桌椅杯书可实例化'),
 'prompt.env_light_zones':      (4,'env_indoor_pack','light zone falloff inverse square 2D; window light gradient','明暗分区可测亮度差'),
 'prompt.env_floor_material':   (4,'env_indoor_pack','procedural terrazzo floor texture generation','地面材质有颗粒'),
 'scene.wall':                  (4,'—','shape grammar city wall brick generation','城墙砖纹可复现'),
 'scene.water':                 (4,'—','procedural water ripple gerstner wave 2D animation','波纹随时间流动不闪'),

 'prompt.dyn_steam':            (5,'—','particle system steam rising buoyancy noise formula','粒子上升且消散'),
 'prompt.dyn_paper_bend':       (5,'—','paper bend deformation curve animation','纸面弯曲可控'),
 'prompt.dyn_light_spot':       (5,'—','dapple light spot caustics animation noise','光斑缓慢晃动'),
 'prompt.dyn_micro_expression': (5,'expression','micro expression animation timing duration','微表情时长 0.2-0.5s'),

 'prompt.subject_face_shape':   (6,'—','face shape parameterization oval round square procedural','脸型参数可切换'),
 'prompt.subject_hair_strands': (6,'—','procedural hair strand generation curve ponytail','发束跟随头动'),
 'prompt.subject_eyelash':      (6,'face_shape','eyelash drawing procedural eye detail','睫毛在眼上沿'),
 'prompt.subject_cheekbone':    (6,'face_shape','procedural face cheekbone geometry shading','颧骨有明暗'),
 'prompt.subject_cloth_inner':  (6,'—','cloth layering inner outer garment procedural collar hem','内外层可分离'),
 'prompt.subject_cloth_lower':  (6,'cloth_inner','trouser leg separation from leg procedural','裤腿与腿分离'),
 'prompt.subject_posture':      (6,'action_sit','sitting posture leaning forward rig','前倾角度可测'),
 'prompt.subject_temperament':  (6,'micro_expression, beat_timeline','temperament through motion timing pacing','节奏参数可调'),

 'prompt.audio_ambient':        (7,'—','procedural ambience sound synthesis noise','环境音可出'),
 'prompt.audio_foley':          (7,'—','foley sound synthesis procedural impact','动作音对齐动作时刻'),
 'prompt.audio_music':          (7,'—','algorithmic music generation piano background','配乐可出'),

 'char.bezier':                 (8,'—','bezier limb tapered organic torso procedural character','四肢锥形可绘'),
 'char.secondary':              (8,'—','secondary motion damped spring hair cloth follow','二级运动滞后于主体'),
 'char.orient':                 (8,'—','8 direction sprite body angle perspective walking','8 朝向可切'),
 'char.arap':                   (8,'—','ARAP as rigid as possible Sorkine Alexa deformation','形变保刚性'),
 'char.bvh':                    (8,'—','BVH motion capture file parse retarget python','BVH 可解析重定向'),
 'post.lumo':                   (8,'—','Lumo 2D shading normal from silhouette weighted','圆被着色成球'),
 'post.timeline':               (8,'—','python render mp4 from frame sequence ffmpeg','直出 MP4'),
 'post.shadow':                 (8,'—','paper cutout shadow offset alpha silhouette','阴影偏移一致'),
}

rows = []
for g in d['groups']:
    for it in g['items']:
        i = it['id']
        w, dep, sh, ver = M.get(i, (9,'—', it.get('search_hint',''), ''))
        if not sh:
            sh = it.get('search_hint','') or '（待补关键词）'
        rows.append({
            'wave': w, 'wave_name': WAVE.get(w,'W9 未排期'),
            'group': g['name'], 'id': i,
            'name': it['name'].split('：')[0] if '：' in it['name'] else it['name'],
            'status': it['status'],
            'dep': dep, 'search_hint': sh,
            'source': it.get('source','') or it.get('src',''),
            'verify': ver,
            'done': '✔' if it['status']=='IMPL' else '',
        })
rows.sort(key=lambda r: (r['wave'], r['id']))

wb = Workbook(); ws = wb.active; ws.title = '缺口工作表'
hdr = ['批次','模块','项ID','缺口名称','当前状态','依赖','搜索关键词(必搜)','业界方案/出处(搜后填)','验证方式','完成']
ws.append(hdr)
for c in range(1, len(hdr)+1):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='4F6228')
    cell.alignment = Alignment(vertical='center')
COLW = [18,16,30,34,10,22,52,46,30,8]
for i,w in enumerate(COLW,1): ws.column_dimensions[get_column_letter(i)].width = w
for r in rows:
    ws.append([r['wave_name'], r['group'], r['id'], r['name'], r['status'],
               r['dep'], r['search_hint'], r['source'], r['verify'], r['done']])
ws.freeze_panes = 'A2'

# 统计页
ws2 = wb.create_sheet('批次汇总')
ws2.append(['批次','项数','说明'])
for w,nm in sorted(WAVE.items()):
    n = sum(1 for r in rows if r['wave']==w)
    ws2.append([nm, n, ''])
ws2.append(['W9 未排期', sum(1 for r in rows if r['wave']==9), ''])
ws2.append([])
ws2.append(['流程', '标完成前必须先搜 search_hint，把出处写进「业界方案/出处」列，否则门禁打回'])

wb.save('缺口工作表.xlsx')

# 回填 JSON：wave / dep / search_hint
for g in d['groups']:
    for it in g['items']:
        i = it['id']
        if i in M:
            w, dep, sh, ver = M[i]
            it['wave'] = w; it['depend'] = dep; it['search_hint'] = sh; it['verify'] = ver
json.dump(d, open(REG,'w',encoding='utf-8'), ensure_ascii=False, indent=2)

n_stub = sum(1 for r in rows if r['status']=='STUB')
print('总项', len(rows), '| STUB 缺口', n_stub, '| 已 IMPL', sum(1 for r in rows if r['status']=='IMPL'))
for w,nm in sorted(WAVE.items()):
    print(' ', nm, sum(1 for r in rows if r['wave']==w))
